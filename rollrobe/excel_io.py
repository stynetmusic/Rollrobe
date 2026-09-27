"""Чтение прайса ПГИ и выгрузка сметы или прайса в xlsx.

Колонки ПГИ, с 4-й строки. Первые три строки — шапка AluRoll, подписей колонок нет:
1 код, 2 флаг (мимо), 3 имя, 4 цвет, 5 флаг (мимо), 6 цена с запятой, 7 код единицы 0/2/3/4.
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from rollrobe.db import Database
from rollrobe.money import line_total, parse_price_cell, q2
from rollrobe.paths import find_data_file

# Имя файла и нужно ли сделать его активным, если активного ещё нет.
SEED_FILES = (
    ("ПГИ-06-12-2024.xlsx", True),
    ("ПГИ-16.09.2024.xlsx", False),
    ("ПГИ-16.09.2024-42.xlsx", False),
)

UNIT_CODES = {0, 2, 3, 4}
PGI_MARKER = "AluRoll"


class ImportFormatError(Exception):
    """Файл не прайс ПГИ или его не удалось прочитать."""


@dataclass(frozen=True)
class PgiRow:
    code: str
    name: str
    color: str
    unit: int
    file_price: Decimal


@dataclass
class ReadResult:
    rows: list[PgiRow]
    skipped: int


def read_pgi(path: Path) -> ReadResult:
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFormatError(f"Не удалось открыть файл: {exc}") from exc
    try:
        sheet = workbook.active
        if sheet is None:
            raise ImportFormatError("В книге нет листа.")
        iterator = sheet.iter_rows(min_row=1, max_col=7, values_only=True)
        try:
            header = next(iterator)
        except StopIteration as exc:
            raise ImportFormatError("Файл пустой.") from exc
        marker = "" if not header or header[0] is None else str(header[0]).strip()
        if marker != PGI_MARKER:
            shown = marker or "пусто"
            raise ImportFormatError(
                "Это не прайс ПГИ. В первой ячейке должно быть «AluRoll», "
                f"сейчас: «{shown}»."
            )
        # Строки 2 и 3 — продолжение шапки, данных там нет.
        next(iterator, None)
        next(iterator, None)
        by_code: dict[str, PgiRow] = {}
        skipped = 0
        for raw in iterator:
            parsed, skip = _parse_row(raw)
            if skip:
                skipped += 1
                continue
            if parsed is not None:
                by_code[parsed.code] = parsed
        return ReadResult(list(by_code.values()), skipped)
    finally:
        workbook.close()


def import_workbook(db: Database, path: Path, *, make_active: bool = False) -> str:
    """Импорт или повторный импорт. Своя цена остаётся, если её меняли."""
    result = read_pgi(path)
    if not result.rows:
        raise ImportFormatError("В файле нет позиций прайса.")
    filename = path.name
    existing = db.find_list(filename)
    if existing is None:
        list_id = db.create_list(path.stem, filename)
        inserted, updated = db.upsert_items(list_id, result.rows)
        verb = "Импортирован"
    else:
        list_id = existing["id"]
        inserted, updated = db.upsert_items(list_id, result.rows)
        db.touch_list(list_id, filename)
        verb = "Обновлён"
    if make_active:
        db.set_active(list_id)
    kept = ""
    if updated:
        kept = " Цены, которые вы меняли, не затерты."
    return (
        f"{verb} прайс «{path.stem}»: добавлено {inserted}, "
        f"обновлено {updated}, пропущено {result.skipped}.{kept}"
    )


def ensure_seed_prices(db: Database, on_progress=None) -> list[str]:
    """Дочитать три файла ПГИ, если их ещё нет в базе. Уже лежащие не перечитывать."""
    notes: list[str] = []
    had_active = db.active_list() is not None
    known_files = {row["filename"] for row in db.lists()}
    known_names = {row["name"] for row in db.lists()}
    for filename, _should_be_preferred in SEED_FILES:
        stem = Path(filename).stem
        if filename in known_files or stem in known_names:
            continue
        path = find_data_file(filename)
        if path is None:
            notes.append(f"Не найден файл {filename}")
            continue
        if on_progress:
            on_progress(f"Импорт {filename}…")
        notes.append(import_workbook(db, path, make_active=False))
        known_files.add(filename)
        known_names.add(stem)
    if not had_active:
        chosen = db.find_list(SEED_FILES[0][0])
        if chosen is None:
            lists = db.lists()
            chosen = lists[0] if lists else None
        if chosen is not None:
            db.set_active(chosen["id"])
    return notes


def export_estimate(path: Path, header: dict, lines: list[dict], unit_labels: dict[int, str]) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Смета"
    fields = (
        ("Номер", header.get("number", "")),
        ("Дата", header.get("date", "")),
        ("Клиент", header.get("client", "")),
        ("Объект", header.get("object", "")),
        ("Ширина, мм", header.get("width", "")),
        ("Высота, мм", header.get("height", "")),
        ("Комментарий", header.get("comment", "")),
    )
    for index, (label, value) in enumerate(fields, start=1):
        sheet.cell(index, 1, label).font = Font(bold=True)
        sheet.cell(index, 2, value)
    sheet.cell(9, 1, "Ширина и высота — поля заказа. В итог не входят.")
    headers = ("Код", "Имя", "Цвет", "Ед.", "Количество", "Цена", "Сумма")
    for column, title in enumerate(headers, start=1):
        sheet.cell(11, column, title).font = Font(bold=True)
    for offset, line in enumerate(lines):
        row = 12 + offset
        amount = line_total(line["qty"], line["price"])
        sheet.cell(row, 1, line["code"])
        sheet.cell(row, 2, line["name"])
        sheet.cell(row, 3, line["color"])
        sheet.cell(row, 4, unit_labels.get(int(line["unit"]), f"ед. {line['unit']}"))
        _money_cell(sheet.cell(row, 5), line["qty"])
        _money_cell(sheet.cell(row, 6), line["price"])
        _money_cell(sheet.cell(row, 7), amount)
    total_row = 13 + len(lines)
    total = Decimal("0.00")
    for line in lines:
        total += line_total(line["qty"], line["price"])
    sheet.cell(total_row, 6, "Итого").font = Font(bold=True)
    total_cell = sheet.cell(total_row, 7, float(q2(total)))
    total_cell.number_format = "0.00"
    total_cell.font = Font(bold=True)
    sheet.freeze_panes = "A12"
    for column, width in enumerate((18, 36, 12, 14, 14, 14, 16), start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    workbook.save(path)


def export_price(path: Path, list_name: str, items: list[dict], unit_labels: dict[int, str]) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Прайс"
    sheet.cell(1, 1, "Прайс").font = Font(bold=True)
    sheet.cell(1, 2, list_name)
    headers = ("Код", "Имя", "Цвет", "Ед.", "Цена файла", "Моя цена", "Изменена")
    for column, title in enumerate(headers, start=1):
        sheet.cell(3, column, title).font = Font(bold=True)
    ordered = sorted(items, key=lambda item: (item["name"].casefold(), item["color"].casefold(), item["code"]))
    for offset, item in enumerate(ordered):
        row = 4 + offset
        sheet.cell(row, 1, item["code"])
        sheet.cell(row, 2, item["name"])
        sheet.cell(row, 3, item["color"])
        sheet.cell(row, 4, unit_labels.get(int(item["unit"]), f"ед. {item['unit']}"))
        _money_cell(sheet.cell(row, 5), item["file_price"])
        _money_cell(sheet.cell(row, 6), item["user_price"])
        sheet.cell(row, 7, "да" if item["user_edited"] else "нет")
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:G{3 + max(len(ordered), 1)}"
    for column, width in enumerate((16, 36, 12, 14, 16, 16, 12), start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    workbook.save(path)


def _money_cell(cell, value: Decimal) -> None:
    cell.value = float(q2(value))
    cell.number_format = "0.00"


def _parse_row(raw) -> tuple[PgiRow | None, bool]:
    cells = list(raw or ())
    if len(cells) < 7:
        cells.extend([None] * (7 - len(cells)))
    cells = cells[:7]
    if _blank_row(cells):
        return None, False
    code = _token(cells[0])
    if not code:
        return None, True
    unit = _unit(cells[6])
    if unit is None:
        return None, True
    try:
        price = parse_price_cell(cells[5])
    except Exception:
        return None, True
    return PgiRow(code, _name(cells[2]), _color(cells[3]), unit, price), False


def _blank_row(cells) -> bool:
    for value in cells:
        if value is None or value is False:
            continue
        if isinstance(value, str) and not value.strip():
            continue
        return False
    return True


def _token(value) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value != value or value in (float("inf"), float("-inf")):
            return None
        if value.is_integer():
            return str(int(value))
        return format(value, "f").rstrip("0").rstrip(".")
    text = str(value).strip()
    return text or None


def _name(value) -> str:
    # В ПГИ пустое имя иногда записано как False.
    if value is None or isinstance(value, bool):
        return ""
    return _token(value) or ""


def _color(value) -> str:
    if value is None or isinstance(value, bool):
        return ""
    return _token(value) or ""


def _unit(value) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int) and value in UNIT_CODES:
        return value
    if isinstance(value, float) and value in UNIT_CODES:
        return int(value)
    if isinstance(value, str) and value.strip() in {"0", "2", "3", "4"}:
        return int(value.strip())
    return None
